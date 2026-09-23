"""无模型端到端冒烟测试：OCR_BACKEND=mock + mock LLM 服务。

覆盖：首页渲染、PDF 上传分析（同步 /analyze 与前端在用的异步 提交+轮询）、
进度回调、结果内容、xlsx 下载与一次性失效。
运行：python tests/test_pipeline_mock.py
"""
import io
import subprocess
import sys
import threading
import time
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

import os  # noqa: E402

os.environ["OCR_BACKEND"] = "mock"
os.environ["LLM_BASE_URL"] = "http://127.0.0.1:9874/v1"
os.environ["MAC_OCR_PORT"] = "9875"

import httpx  # noqa: E402
import uvicorn  # noqa: E402

from scripts.mock_model_server import app as mock_llm  # noqa: E402


def _serve(app, port):
    server = uvicorn.Server(uvicorn.Config(
        app, host="127.0.0.1", port=port, log_level="warning"))
    threading.Thread(target=server.run, daemon=True).start()


_serve(mock_llm, 9874)

from app.main_app import start_server  # noqa: E402

start_server()
base = "http://127.0.0.1:9875"

# 1. 首页渲染（默认提示词已注入）
r = httpx.get(base + "/")
assert r.status_code == 200 and "任务：从给定的商业装箱单" in r.text, "首页渲染异常"
print("[1/5] GET / ok")

# 2. 生成测试 PDF
subprocess.run([sys.executable, str(ROOT / "scripts" / "gen_test_pdf.py")],
               cwd=ROOT, check=True)
pdf = (ROOT / "uploads" / "test_packing_list.pdf").read_bytes()
print("[2/5] test PDF generated")

# 3. 分析（同步接口）
from app.default_prompt import DEFAULT_PROMPT  # noqa: E402

r = httpx.post(base + "/analyze",
               data={"prompt": DEFAULT_PROMPT},
               files={"files": ("test_packing_list.pdf", pdf, "application/pdf")},
               timeout=120)
assert r.status_code == 200, f"{r.status_code}: {r.text[:300]}"
j = r.json()
assert "TOTAL" in j["result"] and j["pages"][0]["ocr"].count("ARSER-PLY") >= 3
assert j["xlsx_url"], "未生成 xlsx"
print("[3/5] /analyze ok ->", j["xlsx_url"])

# 4. xlsx 下载 + 一次性
r = httpx.get(base + j["xlsx_url"])
assert r.status_code == 200
from openpyxl import load_workbook  # noqa: E402

ws = load_workbook(io.BytesIO(r.content)).active
assert ws.max_row >= 8
assert abs(ws.cell(row=ws.max_row, column=6).value - 40953.74) < 0.01
assert httpx.get(base + j["xlsx_url"]).status_code == 404, "xlsx 应一次性"
print("[4/5] xlsx ok")

# 5. 异步任务路径（前端实际用的接口：提交 + 轮询）与进度回调
from app import pipeline  # noqa: E402

msgs = []
pipeline.analyze([("t.pdf", pdf, "application/pdf")], DEFAULT_PROMPT, None,
                 progress=msgs.append)
assert any("OCR 第 1/" in m for m in msgs), f"缺少 OCR 进度：{msgs}"
assert any("汇总" in m for m in msgs), f"缺少 LLM 进度：{msgs}"

r = httpx.post(base + "/analyze/start", data={"prompt": DEFAULT_PROMPT},
               files={"files": ("test_packing_list.pdf", pdf, "application/pdf")},
               timeout=30)
assert r.status_code == 200, f"{r.status_code}: {r.text[:300]}"
jid = r.json()["job_id"]
for _ in range(120):
    pj = httpx.get(base + f"/analyze/status/{jid}", timeout=30).json()
    if pj["status"] != "running":
        break
    time.sleep(0.5)
assert pj["status"] == "done", pj
assert "TOTAL" in pj["result"] and pj["xlsx_url"], "异步任务结果不完整"
assert httpx.get(base + "/analyze/status/deadbeef").status_code == 404, "未知任务应 404"
print(f"[5/5] /analyze/start + /analyze/status ok（进度：{' / '.join(msgs)}）")
print("ALL PASSED")
