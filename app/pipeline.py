"""LLM（ollama）调用与流水线：PDF 转图、两段式分析、Excel 导出。"""
import io
import re
import time
import uuid

import fitz  # PyMuPDF
import httpx
from fastapi import HTTPException

import app.config as cfg
from app.logutil import log


def _log(msg: str) -> None:
    log("pipe", msg)


# ---------- PDF / 图片处理 ----------

def _guess_image_mime(data: bytes, fallback: str) -> str:
    if data.startswith(b"\x89PNG"):
        return "image/png"
    if data.startswith(b"\xff\xd8"):
        return "image/jpeg"
    return fallback if fallback.startswith("image/") else "image/png"


def file_to_page_images(data: bytes, filename: str, content_type: str) -> list[tuple[bytes, str]]:
    """返回 [(png_bytes, mime), ...]，PDF 每页一张图。"""
    lower = (filename or "").lower()
    if lower.endswith(".pdf") or content_type == "application/pdf":
        doc = fitz.open(stream=data, filetype="pdf")
        if doc.page_count > cfg.MAX_PDF_PAGES:
            raise HTTPException(400, f"PDF 页数超过上限（{doc.page_count} 页 > {cfg.MAX_PDF_PAGES} 页）")
        pages = []
        for page in doc:
            pix = page.get_pixmap(matrix=fitz.Matrix(2.0, 2.0))  # ~144dpi
            pages.append((pix.tobytes("png"), "image/png"))
        doc.close()
        return pages
    return [(data, _guess_image_mime(data, content_type))]


# ---------- LLM 调用 ----------

def chat_llm(prompt: str, max_tokens: int) -> str:
    payload = {
        "model": cfg.LLM_MODEL,
        "messages": [{"role": "user", "content": [{"type": "text", "text": prompt}]}],
        "temperature": 0.0,
        "max_tokens": max_tokens,
    }
    r = httpx.post(f"{cfg.llm_base_url()}/chat/completions", json=payload,
                   timeout=cfg.CALL_TIMEOUT_S)
    if r.status_code != 200:
        raise HTTPException(502, f"LLM 服务返回错误：{r.status_code} {r.text[:300]}")
    return r.json()["choices"][0]["message"]["content"]


def apply_prompt(prompt: str, ocr_text: str) -> str:
    text = (
        f"{prompt}\n\n---\n\n以下是由 OCR 从原单逐页识别出的全部文字"
        f"（多页时以 --- PAGE n --- 分隔），请据此完成上面的任务：\n\n{ocr_text}"
    )
    return chat_llm(text, max_tokens=8192)


# ---------- Excel 导出 ----------

def _num(s: str):
    s = s.strip().replace(",", "").replace("$", "").replace("¥", "")
    if not s or not re.fullmatch(r"-?\d+(\.\d+)?", s):
        return None
    return float(s)


def build_xlsx(md_table: str) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font

    # 模型可能输出多轮草稿表格，只取最后一个表格块（修正后的最终结果）
    blocks: list[list[str]] = []
    cur: list[str] | None = None
    for line in md_table.splitlines():
        line = line.strip()
        if "|" in line:
            if cur is None:
                cur = []
                blocks.append(cur)
            cur.append(line)
        else:
            cur = None
    lines = blocks[-1] if blocks else []

    rows = []
    for line in lines:
        # 模型输出可能缺行首/行尾的 |，补齐后再按 | 切分
        if not line.startswith("|"):
            line = "|" + line
        if not line.endswith("|"):
            line += "|"
        cells = [c.strip() for c in line.strip("|").split("|")]
        if all(re.fullmatch(r":?-+:?", c) for c in cells):  # 分隔行
            continue
        rows.append(cells)
    if not rows:
        return b""

    wb = Workbook()
    ws = wb.active
    ws.title = "货物明细"
    money_re = re.compile(r"[$¥]")
    for i, cells in enumerate(rows):
        for j, cell in enumerate(cells[:6]):
            c = ws.cell(row=i + 1, column=j + 1, value=cell)
            c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
            if i > 0:  # 数据行
                v = _num(cell)
                if v is not None and (money_re.search(cell) or j >= 4):
                    c.value = v
                    c.number_format = '"$"#,##0.00'
                elif v is not None and j == 1:
                    c.value = v
                    c.number_format = "0.00"
        if i == 0 or rows[i][0].upper().startswith("TOTAL"):
            for j in range(min(6, len(cells))):
                ws.cell(row=i + 1, column=j + 1).font = Font(bold=True)
    widths = [46, 12, 4, 4, 12, 16]
    for j, w in enumerate(widths):
        ws.column_dimensions[ws.cell(row=1, column=j + 1).column_letter].width = w

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------- 完整流水线 ----------

_xlsx_store: dict[str, bytes] = {}


def analyze(files: list[tuple[str, bytes, str]], prompt: str,
            ocr_fn, progress=None) -> dict:
    """files: [(filename, bytes, content_type), ...]；ocr_fn(png_bytes)->str

    progress(msg)：可选的进度回调，逐阶段上报人话（前端轮询展示用）。
    """
    def note(msg: str) -> None:
        if progress is not None:
            try:
                progress(msg)
            except Exception:  # 进度上报失败不该影响主流程
                pass

    from app.ocr_local import ocr as local_ocr

    if not prompt.strip():
        raise HTTPException(400, "提示词不能为空")

    started = time.time()
    note("正在解析文件（PDF 转页面图）…")
    page_images: list[tuple[str, bytes, str]] = []
    for filename, data, content_type in files:
        if not data:
            continue
        page_images.extend(
            (filename, *t) for t in file_to_page_images(data, filename, content_type))
    if not page_images:
        raise HTTPException(400, "未收到有效的 PDF 或图片文件")
    _log(f"待识别 {len(page_images)} 页，开始逐页 OCR")

    pages_text, per_page = [], []
    for idx, (src, img, _mime) in enumerate(page_images, 1):
        # 首屏提示里点明"首次要加载模型"，避免冷启动那几十秒被当成卡死
        note(f"OCR 第 {idx}/{len(page_images)} 页…" +
             ("（首次运行需加载模型，约 10-60 秒）" if idx == 1 else ""))
        t0 = time.time()
        text = local_ocr.ocr(img) if ocr_fn is None else ocr_fn(img)
        _log(f"OCR 第 {idx}/{len(page_images)} 页完成：用时 {time.time() - t0:.1f}s，"
             f"{len(text)} 字符")
        pages_text.append(f"--- PAGE {idx} ---\n{text}")
        per_page.append({"page": idx, "source": src, "ocr": text})

    note("调用本地模型汇总货物明细…")
    t0 = time.time()
    result = apply_prompt(prompt, "\n\n".join(pages_text))
    _log(f"LLM 汇总完成：用时 {time.time() - t0:.1f}s")
    note("生成 Excel…")

    xlsx_bytes = build_xlsx(result)
    xlsx_url = None
    if xlsx_bytes:
        fid = uuid.uuid4().hex[:12]
        _xlsx_store[fid] = xlsx_bytes
        xlsx_url = f"/download/{fid}/goods.xlsx"

    _log(f"全部完成：总用时 {time.time() - started:.1f}s")
    return {"result": result, "pages": per_page, "xlsx_url": xlsx_url}


def pop_xlsx(fid: str) -> bytes | None:
    return _xlsx_store.pop(fid, None)
