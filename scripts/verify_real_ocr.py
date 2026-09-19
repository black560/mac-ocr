"""真实 OCR（transformers，本机 GPU/CPU）+ mock LLM 的端到端验证。

需已下载模型（scripts/download_models.sh）；不依赖 ollama，LLM 用 mock 服务。
运行：.venv/Scripts/python scripts/verify_real_ocr.py
"""
import sys
import threading
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
os.environ["OCR_BACKEND"] = "transformers"
os.environ["LLM_BASE_URL"] = "http://127.0.0.1:9876/v1"  # mock LLM
os.environ["MAC_OCR_PORT"] = "9877"

import httpx, uvicorn
from scripts.mock_model_server import app as mock_llm

def _serve(app, port):
    server = uvicorn.Server(uvicorn.Config(app, host="127.0.0.1", port=port, log_level="warning"))
    threading.Thread(target=server.run, daemon=True).start()

_serve(mock_llm, 9876)
from app.main_app import start_server
start_server()
base = "http://127.0.0.1:9877"

r = httpx.get(base + "/")
assert r.status_code == 200 and "/marked.min.js" in r.text, "首页未引用本地 marked"
r = httpx.get(base + "/marked.min.js")
assert r.status_code == 200 and "marked v" in r.text, "marked.min.js 路由异常"
print("[1/3] 静态资源 ok（marked 已本地化）")

pdf = open(Path(__file__).resolve().parent.parent / "uploads" / "test_packing_list.pdf", "rb").read()
from app.default_prompt import DEFAULT_PROMPT
r = httpx.post(base + "/analyze", data={"prompt": DEFAULT_PROMPT},
               files={"files": ("test_packing_list.pdf", pdf, "application/pdf")}, timeout=1800)
assert r.status_code == 200, f"{r.status_code}: {r.text[:500]}"
j = r.json()
ocr_text = j["pages"][0]["ocr"]
print("[2/3] analyze ok，OCR 文本", len(ocr_text), "字符")
print("--- OCR 前 400 字 ---")
print(ocr_text[:400])
print("--- LLM 结果 ---")
print(j["result"][:600])
assert j["xlsx_url"], "未生成 xlsx"

r = httpx.get(base + j["xlsx_url"])
assert r.status_code == 200
from openpyxl import load_workbook
ws = load_workbook(io.BytesIO(r.content)).active
print("[3/3] xlsx ok，共", ws.max_row, "行")
print("ALL PASSED")
